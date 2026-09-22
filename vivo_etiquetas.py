from __future__ import annotations

import json
import os
import sys
import subprocess
import tempfile
import math
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from pypdf import PdfReader, PdfWriter

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def abrir_arquivo(caminho):
    caminho = str(caminho)

    if sys.platform.startswith("win"):
        os.startfile(caminho)

    elif sys.platform == "darwin":
        subprocess.Popen(
            ["open", caminho]
        )

    else:
        subprocess.Popen(
            ["xdg-open", caminho]
        )


# ============================================================
# CONFIGURAÇÕES GERAIS
# ============================================================

APP_DIR = Path.home() / ".vivo_etiquetas"
CONFIG_FILE = APP_DIR / "config.json"

CONFIG_PADRAO = {
    "taxa_cartao_mes": 1.00,
    "plano_vivo": "Controle 12GB",
    "valor_plano_vivo": 57.00,
    "arquivo_excel": "",
    "arquivo_pdf_modelo": "",
}

# Excel
ABA_EXCEL = "SMARTPHONES"

COL_NOME = "E"
COL_PRE_PAGO = "F"
COL_ENTRADA = "AX"

# PDF
ETIQUETAS_POR_PAGINA = 10
CAMPOS_POR_ETIQUETA = 13

TOTAL_CAMPOS_POR_PAGINA = (
    ETIQUETAS_POR_PAGINA
    * CAMPOS_POR_ETIQUETA
)


# ============================================================
# CONFIGURAÇÃO
# ============================================================

def carregar_config():

    APP_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    if not CONFIG_FILE.exists():

        salvar_config(
            CONFIG_PADRAO.copy()
        )

        return CONFIG_PADRAO.copy()

    try:

        with open(
            CONFIG_FILE,
            "r",
            encoding="utf-8"
        ) as arquivo:

            dados = json.load(
                arquivo
            )

        config = CONFIG_PADRAO.copy()
        config.update(dados)

        return config

    except (
        OSError,
        json.JSONDecodeError
    ):

        return CONFIG_PADRAO.copy()


def salvar_config(config):

    APP_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    with open(
        CONFIG_FILE,
        "w",
        encoding="utf-8"
    ) as arquivo:

        json.dump(
            config,
            arquivo,
            ensure_ascii=False,
            indent=4
        )


# ============================================================
# MOEDA / DECIMAL
# ============================================================

def moeda(valor):

    return Decimal(
        str(valor)
    ).quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP
    )


def texto_formatado(valor):

    valor = moeda(valor)

    texto = f"{valor:,.2f}"

    texto = texto.replace(
        ",",
        "X"
    )

    texto = texto.replace(
        ".",
        ","
    )

    texto = texto.replace(
        "X",
        "."
    )

    return texto


def formatar_reais(valor):

    return (
        f"R$ "
        f"{texto_formatado(valor)}"
    )


def decimal_brasileiro(valor):

    if valor is None:
        return Decimal("0.00")

    texto = str(
        valor
    ).strip()

    if not texto:
        return Decimal("0.00")

    texto = texto.replace(
        "R$",
        ""
    ).strip()

    if "," in texto:

        texto = texto.replace(
            ".",
            ""
        )

        texto = texto.replace(
            ",",
            "."
        )

    return moeda(
        texto
    )


# ============================================================
# SEGURO
# ============================================================

SEGURO_FAIXAS = [

    {
        "limite": Decimal("2000.00"),
        "valor": Decimal("16.90")
    },

    {
        "limite": Decimal("3000.00"),
        "valor": Decimal("34.90")
    },

    {
        "limite": Decimal("5000.00"),
        "valor": Decimal("69.90")
    },

    {
        "limite": Decimal("10000.00"),
        "valor": Decimal("129.90")
    },

    {
        "limite": Decimal("15000.00"),
        "valor": Decimal("199.90")
    }
]


def calcular_seguro(
    valor_aparelho
):

    valor_aparelho = moeda(
        valor_aparelho
    )

    for faixa in SEGURO_FAIXAS:

        if valor_aparelho <= faixa["limite"]:

            return moeda(
                faixa["valor"]
            )

    raise ValueError(
        "O valor do aparelho ultrapassa "
        "a maior faixa de seguro."
    )


# ============================================================
# PIX
# ============================================================

def calcular_pix(
    controle_entrada
):

    controle_entrada = moeda(
        controle_entrada
    )

    return moeda(
        controle_entrada
        * Decimal("0.90")
    )


# ============================================================
# PARCELAMENTO PRICE
# ============================================================

def calcular_parcela_price(
    valor,
    quantidade_parcelas,
    taxa_mensal
):

    valor = moeda(
        valor
    )

    taxa = (
        Decimal(
            str(taxa_mensal)
        )
        / Decimal("100")
    )

    if quantidade_parcelas <= 0:

        raise ValueError(
            "Quantidade de parcelas inválida."
        )

    if taxa == 0:

        return moeda(
            valor
            / quantidade_parcelas
        )

    fator = (
        Decimal("1")
        + taxa
    ) ** quantidade_parcelas

    parcela = (
        valor
        * (taxa * fator)
        / (fator - Decimal("1"))
    )

    return moeda(
        parcela
    )


# ============================================================
# APARELHO
# ============================================================

class Aparelho:

    def __init__(
        self,
        modelo,
        pre_pago,
        controle_entrada
    ):

        self.modelo = modelo

        self.pre_pago = moeda(
            pre_pago
        )

        self.controle_entrada = moeda(
            controle_entrada
        )

    @property
    def valor_base_seguro(self):

        return self.controle_entrada


# ============================================================
# CÁLCULO COMPLETO DA ETIQUETA
# ============================================================

def calcular_etiqueta(
    aparelho,
    taxa_cartao,
    plano_vivo,
    valor_plano
):

    # Todos os cálculos do aparelho
    # usam exclusivamente o Controle Entrada.

    valor_pix = calcular_pix(
        aparelho.controle_entrada
    )

    valor_seguro = calcular_seguro(
        aparelho.controle_entrada
    )

    parcela_12 = calcular_parcela_price(
        aparelho.controle_entrada,
        12,
        taxa_cartao
    )

    parcela_21 = calcular_parcela_price(
        aparelho.controle_entrada,
        21,
        taxa_cartao
    )

    return {

        "nome_dispositivo":
            aparelho.modelo,

        "valor_pix":
            valor_pix,

        "valor_pre_pago":
            aparelho.pre_pago,

        "valor_seguro":
            valor_seguro,

        "controle_entrada":
            aparelho.controle_entrada,

        "parcela_12x":
            parcela_12,

        "parcela_21x":
            parcela_21,

        "plano_vivo":
            plano_vivo,

        "valor_plano_vivo":
            moeda(
                valor_plano
            )
    }


# ============================================================
# REPOSITÓRIO EXCEL
# ============================================================

class ExcelRepository:

    def __init__(self):

        self.linhas = []


    def carregar(
        self,
        caminho
    ):

        if load_workbook is None:

            raise RuntimeError(
                "A biblioteca openpyxl não está instalada.\n\n"
                "Execute:\n"
                "pip install openpyxl"
            )

        arquivo = Path(
            caminho
        )

        if not arquivo.exists():

            raise FileNotFoundError(
                "Arquivo Excel não encontrado."
            )

        workbook = load_workbook(
            arquivo,
            data_only=True
        )

        if ABA_EXCEL not in workbook.sheetnames:

            raise ValueError(

                f"A aba '{ABA_EXCEL}' "
                "não foi encontrada.\n\n"

                "Abas disponíveis:\n"

                + "\n".join(
                    workbook.sheetnames
                )
            )

        planilha = workbook[
            ABA_EXCEL
        ]

        self.linhas = []

        for numero_linha in range(
            4,
            planilha.max_row + 1
        ):

            nome = planilha[
                f"{COL_NOME}{numero_linha}"
            ].value

            if nome is None:
                continue

            nome = str(
                nome
            ).strip()

            if not nome:
                continue

            pre_pago = planilha[
                f"{COL_PRE_PAGO}{numero_linha}"
            ].value

            controle_entrada = planilha[
                f"{COL_ENTRADA}{numero_linha}"
            ].value

            self.linhas.append({

                "Nome Comercial":
                    nome,

                "Pré-pago":
                    pre_pago,

                "Controle Entrada":
                    controle_entrada,

                "linha_excel":
                    numero_linha
            })

        return len(
            self.linhas
        )


    def pesquisar_modelo(
        self,
        texto
    ):

        texto = str(
            texto
        ).lower().strip()

        if not texto:

            return self.linhas.copy()

        resultados = []

        for linha in self.linhas:

            nome = str(
                linha.get(
                    "Nome Comercial",
                    ""
                )
            ).lower()

            if texto in nome:

                resultados.append(
                    linha
                )

        return resultados


# ============================================================
# MAPA DOS CAMPOS DA ETIQUETA
# ============================================================

def preencher_campos_etiqueta(
    grupo,
    valores
):

    return {

        # Campo 1 - Pré-pago
        grupo[0]:
            valores["pre_pago"],

        # Campo 2 - Seguro
        grupo[1]:
            "Vivo Seguro Celular",

        # Campo 3 - Valor do seguro
        grupo[2]:
            valores["seguro"],

        # Campo 4
        grupo[3]:
            "10x",

        # Campo 5 - Vazio
        grupo[4]:
            "",

        # Campo 6 - Nome do aparelho
        grupo[5]:
            valores["nome"],

        # Campo 7 - Plano
        grupo[6]:
            valores["plano"],

        # Campo 8 - Valor do plano
        grupo[7]:
            valores["valor_plano"],

        # Campo 9 - CONTROLE ENTRADA
        grupo[8]:
            valores["entrada"],

        # Campo 10 - CONTROLE ENTRADA
        grupo[9]:
            valores["entrada"],

        # Campo 11 - Parcela 21x
        grupo[10]:
            valores["parcela_21"],

        # Campo 12 - Parcela 12x
        grupo[11]:
            valores["parcela_12"],

        # Campo 13 - PIX
        grupo[12]:
            valores["pix"],
    }


# ============================================================
# GERA UMA PÁGINA PREENCHIDA
# ============================================================

def gerar_pagina_preenchida(
    arquivo_modelo,
    etiquetas_da_pagina
):

    reader = PdfReader(
        arquivo_modelo
    )

    campos_dict = reader.get_fields()

    if not campos_dict:

        raise ValueError(
            "O PDF modelo não possui "
            "campos preenchíveis."
        )

    campos = list(
        campos_dict.keys()
    )

    if len(campos) != TOTAL_CAMPOS_POR_PAGINA:

        raise ValueError(

            "O PDF modelo não possui "
            "a quantidade esperada de campos.\n\n"

            f"Encontrados: {len(campos)}\n"

            f"Esperados: "
            f"{TOTAL_CAMPOS_POR_PAGINA}"
        )

    # ========================================================
    # CLONA O DOCUMENTO ORIGINAL
    # ========================================================

    writer = PdfWriter()

    writer.clone_document_from_reader(
        reader
    )

    pagina = writer.pages[0]

    # ========================================================
    # PREENCHER AS 10 POSIÇÕES
    # ========================================================

    for numero_etiqueta in range(
        ETIQUETAS_POR_PAGINA
    ):

        etiqueta = (
            etiquetas_da_pagina[
                numero_etiqueta
            ]
        )

        # Posição vazia
        if etiqueta is None:
            continue

        inicio = (
            numero_etiqueta
            * CAMPOS_POR_ETIQUETA
        )

        grupo = campos[
            inicio:
            inicio + CAMPOS_POR_ETIQUETA
        ]

        if len(grupo) != CAMPOS_POR_ETIQUETA:

            raise ValueError(
                f"Grupo da etiqueta "
                f"{numero_etiqueta + 1} inválido."
            )

        valores = {

            "nome":
                etiqueta[
                    "nome_dispositivo"
                ],

            "pre_pago":
                texto_formatado(
                    etiqueta[
                        "valor_pre_pago"
                    ]
                ),

            "seguro":
                texto_formatado(
                    etiqueta[
                        "valor_seguro"
                    ]
                ),

            "entrada":
                texto_formatado(
                    etiqueta[
                        "controle_entrada"
                    ]
                ),

            "parcela_12":
                texto_formatado(
                    etiqueta[
                        "parcela_12x"
                    ]
                ),

            "parcela_21":
                texto_formatado(
                    etiqueta[
                        "parcela_21x"
                    ]
                ),

            "pix":
                texto_formatado(
                    etiqueta[
                        "valor_pix"
                    ]
                ),

            "plano":
                etiqueta[
                    "plano_vivo"
                ],

            "valor_plano":
                texto_formatado(
                    etiqueta[
                        "valor_plano_vivo"
                    ]
                ),
        }

        dados = preencher_campos_etiqueta(
            grupo,
            valores
        )

        writer.update_page_form_field_values(

            pagina,

            dados,

            auto_regenerate=True
        )

    # ========================================================
    # SALVAR PÁGINA TEMPORÁRIA
    # ========================================================

    arquivo_temporario = (
        tempfile.NamedTemporaryFile(

            prefix="vivo_pagina_",

            suffix=".pdf",

            delete=False
        )
    )

    caminho = Path(
        arquivo_temporario.name
    )

    arquivo_temporario.close()

    with open(
        caminho,
        "wb"
    ) as arquivo:

        writer.write(
            arquivo
        )

    return caminho


# ============================================================
# GERA PDF COMPLETO
# ============================================================

def gerar_pdf_completo(
    arquivo_modelo,
    arquivo_saida,
    etiquetas
):

    if not etiquetas:

        raise ValueError(
            "Nenhuma etiqueta foi adicionada."
        )

    quantidade = len(
        etiquetas
    )

    quantidade_paginas = math.ceil(
        quantidade
        / ETIQUETAS_POR_PAGINA
    )

    writer_final = PdfWriter()

    arquivos_temporarios = []

    try:

        for numero_pagina in range(
            quantidade_paginas
        ):

            inicio = (
                numero_pagina
                * ETIQUETAS_POR_PAGINA
            )

            fim = min(
                inicio
                + ETIQUETAS_POR_PAGINA,
                quantidade
            )

            etiquetas_da_pagina = (
                etiquetas[
                    inicio:fim
                ]
            )

            # Completa a página com posições vazias
            while len(
                etiquetas_da_pagina
            ) < ETIQUETAS_POR_PAGINA:

                etiquetas_da_pagina.append(
                    None
                )

            caminho_pagina = (
                gerar_pagina_preenchida(

                    arquivo_modelo,

                    etiquetas_da_pagina
                )
            )

            arquivos_temporarios.append(
                caminho_pagina
            )

            reader_pagina = PdfReader(
                caminho_pagina
            )

            writer_final.add_page(
                reader_pagina.pages[0]
            )

        with open(
            arquivo_saida,
            "wb"
        ) as arquivo:

            writer_final.write(
                arquivo
            )

    finally:

        for caminho in arquivos_temporarios:

            try:

                caminho.unlink(
                    missing_ok=True
                )

            except OSError:
                pass


# ============================================================
# APLICAÇÃO
# ============================================================

class Aplicacao:

    def __init__(
        self,
        root
    ):

        self.root = root

        self.root.title(
            "Vivo - Gerador de Etiquetas"
        )

        self.root.geometry(
            "1100x800"
        )

        self.root.minsize(
            850,
            600
        )

        self.config = (
            carregar_config()
        )

        self.excel = (
            ExcelRepository()
        )

        self.etiquetas = []

        self.resultados = []

        self.indice_etiqueta_selecionada = 0

        self.criar_interface()

        self.carregar_caminhos_salvos()


    # ========================================================
    # INTERFACE
    # ========================================================

    def criar_interface(
        self
    ):

        container = ttk.Frame(
            self.root
        )

        container.pack(
            fill="both",
            expand=True
        )

        self.canvas = tk.Canvas(
            container,
            highlightthickness=0
        )

        self.canvas.pack(
            side="left",
            fill="both",
            expand=True
        )

        scrollbar = ttk.Scrollbar(
            container,
            orient="vertical",
            command=self.canvas.yview
        )

        scrollbar.pack(
            side="right",
            fill="y"
        )

        self.canvas.configure(
            yscrollcommand=scrollbar.set
        )

        self.frame_conteudo = ttk.Frame(
            self.canvas
        )

        self.window_canvas = (
            self.canvas.create_window(
                (0, 0),
                window=self.frame_conteudo,
                anchor="nw"
            )
        )

        self.frame_conteudo.bind(
            "<Configure>",
            self.atualizar_scroll
        )

        self.canvas.bind(
            "<Configure>",
            self.ajustar_largura_canvas
        )

        self.canvas.bind_all(
            "<MouseWheel>",
            self.rolar_mouse
        )

        self.canvas.bind_all(
            "<Button-4>",
            self.rolar_mouse_linux
        )

        self.canvas.bind_all(
            "<Button-5>",
            self.rolar_mouse_linux
        )

        self.montar_conteudo()


    def montar_conteudo(
        self
    ):

        frame = self.frame_conteudo

        # ====================================================
        # TÍTULO
        # ====================================================

        titulo = ttk.Label(
            frame,
            text="Gerador de Etiquetas Vivo",
            font=(
                "TkDefaultFont",
                18,
                "bold"
            )
        )

        titulo.pack(
            pady=(15, 5)
        )

        subtitulo = ttk.Label(
            frame,
            text=(
                "Excel → cálculo automático "
                "→ etiquetas → PDF"
            )
        )

        subtitulo.pack(
            pady=(0, 15)
        )

        # ====================================================
        # ARQUIVOS
        # ====================================================

        frame_arquivos = ttk.LabelFrame(
            frame,
            text="Arquivos"
        )

        frame_arquivos.pack(
            fill="x",
            padx=15,
            pady=5
        )

        frame_arquivos.columnconfigure(
            1,
            weight=1
        )

        ttk.Label(
            frame_arquivos,
            text="Excel:"
        ).grid(
            row=0,
            column=0,
            padx=8,
            pady=8
        )

        self.var_excel = tk.StringVar()

        ttk.Entry(
            frame_arquivos,
            textvariable=self.var_excel
        ).grid(
            row=0,
            column=1,
            sticky="ew",
            padx=5
        )

        ttk.Button(
            frame_arquivos,
            text="Selecionar",
            command=self.selecionar_excel
        ).grid(
            row=0,
            column=2,
            padx=5
        )

        ttk.Button(
            frame_arquivos,
            text="Carregar Excel",
            command=self.carregar_excel
        ).grid(
            row=0,
            column=3,
            padx=5
        )

        ttk.Label(
            frame_arquivos,
            text="PDF modelo:"
        ).grid(
            row=1,
            column=0,
            padx=8,
            pady=8
        )

        self.var_pdf = tk.StringVar()

        ttk.Entry(
            frame_arquivos,
            textvariable=self.var_pdf
        ).grid(
            row=1,
            column=1,
            sticky="ew",
            padx=5
        )

        ttk.Button(
            frame_arquivos,
            text="Selecionar",
            command=self.selecionar_pdf
        ).grid(
            row=1,
            column=2,
            padx=5
        )

        # ====================================================
        # CONFIGURAÇÕES
        # ====================================================

        frame_config = ttk.LabelFrame(
            frame,
            text="Configurações"
        )

        frame_config.pack(
            fill="x",
            padx=15,
            pady=5
        )

        ttk.Label(
            frame_config,
            text="Taxa cartão (% a.m.):"
        ).grid(
            row=0,
            column=0,
            padx=8,
            pady=8
        )

        self.var_taxa = tk.StringVar(
            value=str(
                self.config[
                    "taxa_cartao_mes"
                ]
            )
        )

        ttk.Entry(
            frame_config,
            textvariable=self.var_taxa,
            width=10
        ).grid(
            row=0,
            column=1,
            padx=5
        )

        ttk.Label(
            frame_config,
            text="Plano:"
        ).grid(
            row=0,
            column=2,
            padx=8
        )

        self.var_plano = tk.StringVar(
            value=self.config[
                "plano_vivo"
            ]
        )

        ttk.Entry(
            frame_config,
            textvariable=self.var_plano,
            width=25
        ).grid(
            row=0,
            column=3,
            padx=5
        )

        ttk.Label(
            frame_config,
            text="Valor plano:"
        ).grid(
            row=0,
            column=4,
            padx=8
        )

        self.var_valor_plano = tk.StringVar(
            value=str(
                self.config[
                    "valor_plano_vivo"
                ]
            )
        )

        ttk.Entry(
            frame_config,
            textvariable=self.var_valor_plano,
            width=10
        ).grid(
            row=0,
            column=5,
            padx=5
        )

        # ====================================================
        # PESQUISA
        # ====================================================

        frame_pesquisa = ttk.LabelFrame(
            frame,
            text="Pesquisar aparelho"
        )

        frame_pesquisa.pack(
            fill="x",
            padx=15,
            pady=5
        )

        self.var_busca = tk.StringVar()

        entrada_busca = ttk.Entry(
            frame_pesquisa,
            textvariable=self.var_busca
        )

        entrada_busca.pack(
            side="left",
            fill="x",
            expand=True,
            padx=8,
            pady=8
        )

        entrada_busca.bind(
            "<KeyRelease>",
            lambda event: self.pesquisar()
        )

        ttk.Button(
            frame_pesquisa,
            text="Pesquisar",
            command=self.pesquisar
        ).pack(
            side="left",
            padx=5
        )

        # ====================================================
        # RESULTADOS
        # ====================================================

        frame_resultados = ttk.Frame(
            frame
        )

        frame_resultados.pack(
            fill="both",
            expand=True,
            padx=15,
            pady=5
        )

        frame_resultados.columnconfigure(
            0,
            weight=1
        )

        frame_resultados.columnconfigure(
            1,
            weight=1
        )

        frame_lista = ttk.LabelFrame(
            frame_resultados,
            text="Aparelhos encontrados"
        )

        frame_lista.grid(
            row=0,
            column=0,
            sticky="nsew",
            padx=(0, 5)
        )

        self.lista_modelos = tk.Listbox(
            frame_lista,
            height=10
        )

        self.lista_modelos.pack(
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        self.lista_modelos.bind(
            "<<ListboxSelect>>",
            self.selecionar_modelo
        )

        # ====================================================
        # DADOS DO APARELHO
        # ====================================================

        frame_dados = ttk.LabelFrame(
            frame_resultados,
            text="Dados do aparelho"
        )

        frame_dados.grid(
            row=0,
            column=1,
            sticky="nsew",
            padx=(5, 0)
        )

        frame_dados.columnconfigure(
            1,
            weight=1
        )

        self.var_nome = tk.StringVar()
        self.var_pre_pago = tk.StringVar()
        self.var_entrada = tk.StringVar()

        campos = [

            (
                "Nome:",
                self.var_nome
            ),

            (
                "Pré-pago:",
                self.var_pre_pago
            ),

            (
                "Controle Entrada:",
                self.var_entrada
            )
        ]

        for linha, (
            nome,
            variavel
        ) in enumerate(campos):

            ttk.Label(
                frame_dados,
                text=nome
            ).grid(
                row=linha,
                column=0,
                sticky="w",
                padx=8,
                pady=8
            )

            ttk.Entry(
                frame_dados,
                textvariable=variavel,
                state="readonly"
            ).grid(
                row=linha,
                column=1,
                sticky="ew",
                padx=8,
                pady=8
            )

        ttk.Button(
            frame_dados,
            text="Calcular aparelho",
            command=self.calcular_atual
        ).grid(
            row=3,
            column=0,
            columnspan=2,
            pady=12
        )

        # ====================================================
        # ETIQUETAS
        # ====================================================

        frame_etiquetas = ttk.LabelFrame(
            frame,
            text="Etiquetas adicionadas"
        )

        frame_etiquetas.pack(
            fill="x",
            padx=15,
            pady=5
        )

        self.lista_etiquetas = tk.Listbox(
            frame_etiquetas,
            height=10
        )

        self.lista_etiquetas.pack(
            side="left",
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        self.lista_etiquetas.bind(
            "<<ListboxSelect>>",
            self.selecionar_etiqueta
        )

        frame_botoes = ttk.Frame(
            frame_etiquetas
        )

        frame_botoes.pack(
            side="right",
            padx=8
        )

        ttk.Button(
            frame_botoes,
            text="➕ Adicionar etiqueta",
            command=self.adicionar_etiqueta
        ).pack(
            fill="x",
            pady=3
        )

        ttk.Button(
            frame_botoes,
            text="🗑 Remover selecionada",
            command=self.remover_etiqueta
        ).pack(
            fill="x",
            pady=3
        )

        self.var_status = tk.StringVar(
            value="Nenhuma etiqueta adicionada."
        )

        ttk.Label(
            frame,
            textvariable=self.var_status,
            font=(
                "TkDefaultFont",
                10,
                "bold"
            )
        ).pack(
            pady=5
        )

        # ====================================================
        # BOTÕES PDF
        # ====================================================

        frame_pdf = ttk.Frame(
            frame
        )

        frame_pdf.pack(
            pady=(5, 20)
        )

        ttk.Button(
            frame_pdf,
            text="👁 Pré-visualizar PDF",
            command=self.previsualizar_pdf
        ).pack(
            side="left",
            padx=8,
            ipadx=12,
            ipady=7
        )

        ttk.Button(
            frame_pdf,
            text="💾 Gerar PDF",
            command=self.salvar_pdf
        ).pack(
            side="left",
            padx=8,
            ipadx=12,
            ipady=7
        )


    # ========================================================
    # SCROLL
    # ========================================================

    def atualizar_scroll(
        self,
        event=None
    ):

        self.canvas.configure(
            scrollregion=self.canvas.bbox(
                "all"
            )
        )


    def ajustar_largura_canvas(
        self,
        event
    ):

        self.canvas.itemconfigure(
            self.window_canvas,
            width=event.width
        )


    def rolar_mouse(
        self,
        event
    ):

        self.canvas.yview_scroll(
            int(
                -1
                * (
                    event.delta
                    / 120
                )
            ),
            "units"
        )


    def rolar_mouse_linux(
        self,
        event
    ):

        if event.num == 4:

            self.canvas.yview_scroll(
                -3,
                "units"
            )

        elif event.num == 5:

            self.canvas.yview_scroll(
                3,
                "units"
            )


    # ========================================================
    # CONFIGURAÇÕES SALVAS
    # ========================================================

    def carregar_caminhos_salvos(
        self
    ):

        self.var_excel.set(
            self.config.get(
                "arquivo_excel",
                ""
            )
        )

        self.var_pdf.set(
            self.config.get(
                "arquivo_pdf_modelo",
                ""
            )
        )


    def salvar_config_atual(
        self
    ):

        try:

            taxa = float(
                self.var_taxa
                .get()
                .replace(
                    ",",
                    "."
                )
            )

            valor_plano = float(
                self.var_valor_plano
                .get()
                .replace(
                    ",",
                    "."
                )
            )

        except ValueError:

            raise ValueError(
                "Taxa do cartão ou "
                "valor do plano inválido."
            )

        self.config[
            "taxa_cartao_mes"
        ] = taxa

        self.config[
            "plano_vivo"
        ] = self.var_plano.get()

        self.config[
            "valor_plano_vivo"
        ] = valor_plano

        self.config[
            "arquivo_excel"
        ] = self.var_excel.get()

        self.config[
            "arquivo_pdf_modelo"
        ] = self.var_pdf.get()

        salvar_config(
            self.config
        )


    # ========================================================
    # SELECIONAR ARQUIVOS
    # ========================================================

    def selecionar_excel(
        self
    ):

        caminho = filedialog.askopenfilename(

            title="Selecionar Excel",

            filetypes=[

                (
                    "Excel",
                    "*.xlsx *.xlsm"
                ),

                (
                    "Todos os arquivos",
                    "*.*"
                )
            ]
        )

        if caminho:

            self.var_excel.set(
                caminho
            )

            self.salvar_config_atual()


    def selecionar_pdf(
        self
    ):

        caminho = filedialog.askopenfilename(

            title="Selecionar PDF modelo",

            filetypes=[

                (
                    "PDF",
                    "*.pdf"
                )
            ]
        )

        if caminho:

            self.var_pdf.set(
                caminho
            )

            self.salvar_config_atual()


    # ========================================================
    # EXCEL
    # ========================================================

    def carregar_excel(
        self
    ):

        caminho = (
            self.var_excel
            .get()
            .strip()
        )

        if not caminho:

            messagebox.showwarning(
                "Excel",
                "Selecione um arquivo Excel primeiro."
            )

            return

        try:

            quantidade = (
                self.excel.carregar(
                    caminho
                )
            )

            self.resultados = (
                self.excel.pesquisar_modelo(
                    ""
                )
            )

            self.atualizar_lista_modelos()

            self.salvar_config_atual()

            messagebox.showinfo(
                "Excel carregado",
                (
                    f"{quantidade} aparelhos "
                    "carregados com sucesso."
                )
            )

        except Exception as erro:

            messagebox.showerror(
                "Erro ao carregar Excel",
                str(erro)
            )


    def pesquisar(
        self
    ):

        texto = (
            self.var_busca.get()
        )

        self.resultados = (
            self.excel.pesquisar_modelo(
                texto
            )
        )

        self.atualizar_lista_modelos()


    def atualizar_lista_modelos(
        self
    ):

        self.lista_modelos.delete(
            0,
            tk.END
        )

        for linha in self.resultados:

            self.lista_modelos.insert(
                tk.END,
                linha[
                    "Nome Comercial"
                ]
            )


    def selecionar_modelo(
        self,
        event=None
    ):

        selecao = (
            self.lista_modelos
            .curselection()
        )

        if not selecao:
            return

        indice = selecao[0]

        linha = (
            self.resultados[
                indice
            ]
        )

        self.preencher_dados(
            linha
        )


    def preencher_dados(
        self,
        linha
    ):

        self.var_nome.set(
            linha.get(
                "Nome Comercial",
                ""
            )
        )

        self.var_pre_pago.set(
            texto_formatado(
                decimal_brasileiro(
                    linha.get(
                        "Pré-pago",
                        0
                    )
                )
            )
        )

        self.var_entrada.set(
            texto_formatado(
                decimal_brasileiro(
                    linha.get(
                        "Controle Entrada",
                        0
                    )
                )
            )
        )


    # ========================================================
    # CÁLCULO
    # ========================================================

    def obter_aparelho_atual(
        self
    ):

        return Aparelho(

            self.var_nome.get(),

            decimal_brasileiro(
                self.var_pre_pago.get()
            ),

            decimal_brasileiro(
                self.var_entrada.get()
            )
        )


    def calcular_atual(
        self
    ):

        if not self.var_nome.get():

            messagebox.showwarning(
                "Aparelho",
                "Selecione um aparelho primeiro."
            )

            return None

        try:

            aparelho = (
                self.obter_aparelho_atual()
            )

            taxa = float(
                self.var_taxa
                .get()
                .replace(
                    ",",
                    "."
                )
            )

            valor_plano = (
                decimal_brasileiro(
                    self.var_valor_plano.get()
                )
            )

            resultado = (
                calcular_etiqueta(

                    aparelho,

                    taxa,

                    self.var_plano.get(),

                    valor_plano
                )
            )

            self.mostrar_resultado(
                resultado
            )

            return resultado

        except Exception as erro:

            messagebox.showerror(
                "Erro no cálculo",
                str(erro)
            )

            return None


    def mostrar_resultado(
        self,
        resultado
    ):

        janela = tk.Toplevel(
            self.root
        )

        janela.title(
            "Resultado do cálculo"
        )

        janela.geometry(
            "450x500"
        )

        texto = tk.Text(
            janela,
            wrap="word"
        )

        texto.pack(
            fill="both",
            expand=True,
            padx=10,
            pady=10
        )

        texto.insert(
            "1.0",

            f"APARELHO\n"
            f"{resultado['nome_dispositivo']}\n\n"

            f"PRÉ-PAGO\n"
            f"{formatar_reais(resultado['valor_pre_pago'])}\n\n"

            f"CONTROLE ENTRADA\n"
            f"{formatar_reais(resultado['controle_entrada'])}\n\n"

            f"SEGURO\n"
            f"{formatar_reais(resultado['valor_seguro'])}\n\n"

            f"PIX\n"
            f"{formatar_reais(resultado['valor_pix'])}\n\n"

            f"12X\n"
            f"{formatar_reais(resultado['parcela_12x'])}\n\n"

            f"21X\n"
            f"{formatar_reais(resultado['parcela_21x'])}\n\n"

            f"PLANO\n"
            f"{resultado['plano_vivo']}\n"

            f"{formatar_reais(resultado['valor_plano_vivo'])}"
        )

        texto.config(
            state="disabled"
        )

        ttk.Button(
            janela,
            text="Fechar",
            command=janela.destroy
        ).pack(
            pady=10
        )


    # ========================================================
    # ETIQUETAS
    # ========================================================

    def selecionar_etiqueta(
        self,
        event=None
    ):

        selecao = (
            self.lista_etiquetas
            .curselection()
        )

        if not selecao:
            return

        self.indice_etiqueta_selecionada = (
            selecao[0]
        )


    def adicionar_etiqueta(
        self
    ):

        resultado = (
            self.calcular_atual()
        )

        if resultado is None:
            return

        self.etiquetas.append(
            resultado
        )

        indice = (
            len(
                self.etiquetas
            ) - 1
        )

        self.atualizar_lista_etiquetas()

        self.lista_etiquetas.selection_clear(
            0,
            tk.END
        )

        self.lista_etiquetas.selection_set(
            indice
        )

        self.lista_etiquetas.see(
            indice
        )

        self.lista_etiquetas.activate(
            indice
        )

        self.indice_etiqueta_selecionada = (
            indice
        )

        self.var_status.set(
            (
                f"Etiqueta {indice + 1} "
                "adicionada: "
                f"{resultado['nome_dispositivo']}."
            )
        )


    def remover_etiqueta(
        self
    ):

        if not self.etiquetas:
            return

        selecao = (
            self.lista_etiquetas
            .curselection()
        )

        if not selecao:

            indice = (
                len(
                    self.etiquetas
                ) - 1
            )

        else:

            indice = selecao[0]

        etiqueta = (
            self.etiquetas[
                indice
            ]
        )

        del self.etiquetas[
            indice
        ]

        self.atualizar_lista_etiquetas()

        if self.etiquetas:

            novo_indice = min(

                indice,

                len(
                    self.etiquetas
                ) - 1
            )

            self.lista_etiquetas.selection_set(
                novo_indice
            )

            self.lista_etiquetas.see(
                novo_indice
            )

            self.indice_etiqueta_selecionada = (
                novo_indice
            )

        else:

            self.indice_etiqueta_selecionada = 0

        self.var_status.set(
            (
                "Etiqueta removida: "
                f"{etiqueta['nome_dispositivo']}."
            )
        )


    def atualizar_lista_etiquetas(
        self
    ):

        self.lista_etiquetas.delete(
            0,
            tk.END
        )

        quantidade = len(
            self.etiquetas
        )

        if quantidade == 0:

            self.lista_etiquetas.insert(
                tk.END,
                "Nenhuma etiqueta adicionada."
            )

        else:

            for indice, etiqueta in enumerate(
                self.etiquetas
            ):

                numero_pagina = (
                    indice
                    // ETIQUETAS_POR_PAGINA
                ) + 1

                posicao_pagina = (
                    indice
                    % ETIQUETAS_POR_PAGINA
                ) + 1

                self.lista_etiquetas.insert(
                    tk.END,

                    (
                        f"{indice + 1:03d}  |  "
                        f"Página {numero_pagina} "
                        f"• posição {posicao_pagina}  |  "
                        f"{etiqueta['nome_dispositivo']}"
                    )
                )

        paginas = 0

        if quantidade:

            paginas = math.ceil(
                quantidade
                / ETIQUETAS_POR_PAGINA
            )

        self.var_status.set(
            (
                f"{quantidade} etiqueta(s) | "
                f"{paginas} página(s) no PDF | "
                f"{ETIQUETAS_POR_PAGINA} por página"
            )
        )


    # ========================================================
    # VALIDAÇÃO PDF
    # ========================================================

    def validar_para_pdf(
        self
    ):

        arquivo_modelo = (
            self.var_pdf
            .get()
            .strip()
        )

        if not arquivo_modelo:

            raise ValueError(
                "Selecione o PDF modelo primeiro."
            )

        caminho = Path(
            arquivo_modelo
        )

        if not caminho.exists():

            raise FileNotFoundError(
                "O arquivo PDF modelo "
                "não foi encontrado."
            )

        if not self.etiquetas:

            raise ValueError(
                "Nenhuma etiqueta foi adicionada."
            )

        return (
            arquivo_modelo,
            len(
                self.etiquetas
            )
        )


    # ========================================================
    # PRÉ-VISUALIZAÇÃO
    # ========================================================

    def previsualizar_pdf(
        self
    ):

        try:

            arquivo_modelo, quantidade = (
                self.validar_para_pdf()
            )

            arquivo_preview = (
                tempfile.NamedTemporaryFile(

                    prefix="vivo_preview_",

                    suffix=".pdf",

                    delete=False
                )
            )

            caminho_preview = Path(
                arquivo_preview.name
            )

            arquivo_preview.close()

            gerar_pdf_completo(

                arquivo_modelo,

                caminho_preview,

                self.etiquetas.copy()
            )

            abrir_arquivo(caminho_preview)

            paginas = math.ceil(
                quantidade
                / ETIQUETAS_POR_PAGINA
            )

            self.var_status.set(
                (
                    "Pré-visualização aberta: "
                    f"{quantidade} etiqueta(s) "
                    f"em {paginas} página(s)."
                )
            )

        except Exception as erro:

            messagebox.showerror(
                "Erro na pré-visualização",
                str(erro)
            )


    # ========================================================
    # SALVAR PDF
    # ========================================================

    def salvar_pdf(
        self
    ):

        try:

            arquivo_modelo, quantidade = (
                self.validar_para_pdf()
            )

            self.salvar_config_atual()

            paginas = math.ceil(
                quantidade
                / ETIQUETAS_POR_PAGINA
            )

            arquivo_saida = (
                filedialog.asksaveasfilename(

                    title="Salvar PDF preenchido",

                    defaultextension=".pdf",

                    initialfile=(
                        "etiquetas_vivo_"
                        "preenchidas.pdf"
                    ),

                    filetypes=[
                        (
                            "PDF",
                            "*.pdf"
                        )
                    ]
                )
            )

            if not arquivo_saida:
                return

            gerar_pdf_completo(

                arquivo_modelo,

                arquivo_saida,

                self.etiquetas.copy()
            )

            resposta = (
                messagebox.askyesno(

                    "PDF gerado",

                    (
                        "PDF gerado com sucesso!\n\n"
                        f"Etiquetas: {quantidade}\n"
                        f"Páginas: {paginas}\n\n"
                        "Deseja abrir o PDF agora?"
                    )
                )
            )

            if resposta:

                abrir_arquivo(
                    arquivo_saida
                )

            self.var_status.set(
                (
                    "PDF gerado: "
                    f"{quantidade} etiqueta(s), "
                    f"{paginas} página(s)."
                )
            )

        except Exception as erro:

            messagebox.showerror(
                "Erro ao gerar PDF",
                str(erro)
            )


# ============================================================
# MAIN
# ============================================================

def main():

    root = tk.Tk()

    Aplicacao(
        root
    )

    root.mainloop()


if __name__ == "__main__":

    main()
